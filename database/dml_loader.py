"""
Cargador DML desde Excel.
Ejecutar después de generar el Excel:
  python data/generar_datos.py
  python database/dml_loader.py
"""
import pathlib
import openpyxl
import sys
import os

# Agregar directorio raíz al path
sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))

from config.db import get_connection, _bootstrap

EXCEL_FILE = pathlib.Path(__file__).parent.parent / "data" / "datos_parqueadero.xlsx"


def load_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = rows[0]
    data = rows[1:]
    return headers, data


def run():
    if not EXCEL_FILE.exists():
        print(f"ERROR: No se encontró {EXCEL_FILE}")
        print("Ejecute primero: python data/generar_datos.py")
        return

    print("Iniciando carga desde Excel...")

    # Bootstrap DB + tablas si no existen
    _bootstrap()
    conn = get_connection()
    cursor = conn.cursor()

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)

        # 1. TipoVehiculo
        _insert_table(cursor, conn, wb, "TipoVehiculo",
                      "REPLACE INTO TipoVehiculo (id_tipo_vehiculo, nombre) VALUES (%s,%s)")

        # 2. Tarifa
        _insert_table(cursor, conn, wb, "Tarifa",
                      "REPLACE INTO Tarifa (id_tarifa, valor_hora, descripcion, id_tipo_vehiculo) VALUES (%s,%s,%s,%s)")

        # 3. Zona
        _insert_table(cursor, conn, wb, "Zona",
                      "REPLACE INTO Zona (id_zona, nombre, piso, tipo_zona, capacidad) VALUES (%s,%s,%s,%s,%s)")

        # 4. EspacioParqueo
        _insert_table(cursor, conn, wb, "EspacioParqueo",
                      "REPLACE INTO EspacioParqueo (id_espacio, numero, estado, tipo_espacio, id_zona) VALUES (%s,%s,%s,%s,%s)")

        # 5. Empresa
        _insert_table(cursor, conn, wb, "Empresa",
                      "REPLACE INTO Empresa (id_empresa, nombre, nit, direccion, telefono, ciudad, estado_convenio) VALUES (%s,%s,%s,%s,%s,%s,%s)")

        # 6. Convenio
        _insert_table(cursor, conn, wb, "Convenio",
                      "REPLACE INTO Convenio (id_convenio, id_empresa, estado, porcentaje_descuento) VALUES (%s,%s,%s,%s)")

        # 7. Usuario
        _insert_table(cursor, conn, wb, "Usuario",
                      "REPLACE INTO Usuario (cedula, primer_nombre, segundo_nombre, primer_apellido, segundo_apellido, correo) VALUES (%s,%s,%s,%s,%s,%s)")

        # 8. Administrador
        _insert_table(cursor, conn, wb, "Administrador",
                      "REPLACE INTO Administrador (cedula, contrasenia, fecha_asignacion, codigo_interno) VALUES (%s,%s,%s,%s)")

        # 9. Operador
        _insert_table(cursor, conn, wb, "Operador",
                      "REPLACE INTO Operador (cedula, codigo_interno, estado) VALUES (%s,%s,%s)")

        # 10. Turno
        _insert_table(cursor, conn, wb, "Turno",
                      "REPLACE INTO Turno (id_turno, id_operador, fecha_inicio_turno, fecha_final_turno) VALUES (%s,%s,%s,%s)")

        # 11. Suscriptor
        _insert_table(cursor, conn, wb, "Suscriptor",
                      "REPLACE INTO Suscriptor (cedula, fecha_registro) VALUES (%s,%s)")

        # 12. Telefono
        _insert_table(cursor, conn, wb, "Telefono",
                      "REPLACE INTO Telefono (cedula_usuario, telefono) VALUES (%s,%s)")

        # 13. Vehiculo
        _insert_table(cursor, conn, wb, "Vehiculo",
                      "REPLACE INTO Vehiculo (id_vehiculo, placa, color, marca, modelo, cedula_usuario, id_tipo_vehiculo, id_empresa) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)")

        # 14. Suscripcion
        _insert_table(cursor, conn, wb, "Suscripcion",
                      "REPLACE INTO Suscripcion (id_suscripcion, fecha_inicio, fecha_final, estado, horario_permitido_inicio, horario_permitido_final, id_vehiculo) VALUES (%s,%s,%s,%s,%s,%s,%s)")

        # 15. SuscriptorSuscripcion
        _insert_table(cursor, conn, wb, "SuscriptorSuscripcion",
                      "REPLACE INTO SuscriptorSuscripcion (id_suscriptor, id_suscripcion) VALUES (%s,%s)")

        # 16. SesionParqueo
        _insert_table(cursor, conn, wb, "SesionParqueo",
                      "REPLACE INTO SesionParqueo (id_sesion, fecha_inicio, fecha_fin, tiempo, id_vehiculo, id_espacio) VALUES (%s,%s,%s,%s,%s,%s)")

        # 17. Factura
        _insert_table(cursor, conn, wb, "Factura",
                      "REPLACE INTO Factura (id_factura, fecha_ingreso, fecha_salida, tiempo, valor_total, estado_pago, descuento_aplicado, tipo_cobro, id_sesion, id_tarifa) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)")

        # 18. Multa
        _insert_table(cursor, conn, wb, "Multa",
                      "REPLACE INTO Multa (id_multa, fecha, motivo, valor, estado, id_sesion) VALUES (%s,%s,%s,%s,%s,%s)")

        # 19. Notificacion
        _insert_table(cursor, conn, wb, "Notificacion",
                      "REPLACE INTO Notificacion (id_notificacion, fecha, mensaje, tipo_notificacion, cedula_usuario) VALUES (%s,%s,%s,%s,%s)")

        print("\nOK Carga completada exitosamente.")

    except Exception as e:
        conn.rollback()
        print(f"ERROR: {e}")
        import traceback
        traceback.print_exc()
    finally:
        cursor.close()
        conn.close()


def _insert_table(cursor, conn, wb, sheet_name, sql):
    if sheet_name not in wb.sheetnames:
        print(f"  SKIP: hoja '{sheet_name}' no encontrada.")
        return
    ws = wb[sheet_name]
    headers, data = load_sheet(ws)
    count = 0
    errors = 0
    for row in data:
        if all(v is None for v in row):
            continue
        try:
            cursor.execute(sql, tuple(row))
            count += 1
        except Exception as e:
            errors += 1
            if errors <= 3:
                print(f"    WARN ({sheet_name}): {e} — fila: {row}")
    conn.commit()
    print(f"  OK {sheet_name}: {count} registros insertados" +
          (f" ({errors} errores)" if errors else ""))


if __name__ == "__main__":
    run()
